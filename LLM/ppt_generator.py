import io
import os
import re
import copy
import urllib.request
from datetime import datetime
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu
import lxml.etree as etree


def safe_str(val, default=""):
    """Convert value to clean string, handling nan/None/empty gracefully."""
    if val is None:
        return default
    s = str(val).strip()
    if s.lower() in ('nan', 'none', 'n/a', 'available on request / document image', ''):
        return default
    return s


def format_currency(val, prefix="INR ", suffix=" / Sq. Ft. / Month"):
    """Format a rent/charge value nicely."""
    s = safe_str(val)
    if not s:
        return "Quote on Request"
    # Already formatted
    if 'INR' in s.upper() or 'included' in s.lower() or 'tbd' in s.lower() or 'quote' in s.lower():
        return s
    try:
        num = float(str(s).replace(',', ''))
        return f"{prefix}{num:,.0f}{suffix}"
    except (ValueError, TypeError):
        return s


def format_months(val, default=""):
    """
    '36.0' -> '36 Months'.

    The caller used to loop over a list of locals reassigning the loop
    variable, which does nothing, so these rendered as bare floats: a client
    saw "Lease Term: 36.0".
    """
    s = safe_str(val, default)
    if not s:
        return default
    try:
        number = float(str(s).replace(",", ""))
    except (ValueError, TypeError):
        return s
    if number <= 0:
        return default
    whole = int(number)
    return "%d Month%s" % (whole, "" if whole == 1 else "s")


def format_percent(val, default=""):
    """
    Escalation arrives from Excel as a fraction: 0.08 is 8%, and rendered raw
    it read as "Rental Escalation: 0.08".
    """
    s = safe_str(val, default)
    if not s:
        return default
    if "%" in s:
        return s
    try:
        number = float(str(s).replace(",", ""))
    except (ValueError, TypeError):
        return s
    if number <= 1:                     # a fraction, not a percentage
        number *= 100
    text = ("%g" % round(number, 2))
    return text + "%"


def format_quantity(area, seats=None):
    """
    Managed office is sold by the seat. Rendering only the area gave
    "Area Offered: 0 Sq. Ft." on an option that had 105 seats available.
    """
    area_text = format_area(area) if safe_str(area) else ""
    if area_text and area_text != "Available on Request" and not area_text.startswith("0 "):
        return area_text
    seat_text = safe_str(seats)
    if seat_text:
        try:
            count = int(float(seat_text.replace(",", "")))
            if count > 0:
                return "%s Seats" % "{:,}".format(count)
        except (ValueError, TypeError):
            pass
    return "Available on Request"


def format_area(val):
    """Format area value nicely."""
    s = safe_str(val)
    if not s:
        return "Available on Request"
    try:
        num = float(str(s).replace(',', ''))
        if num > 0:
            return f"{num:,.0f} Sq. Ft."
    except (ValueError, TypeError):
        pass
    if 'sq' not in s.lower() and 'sft' not in s.lower():
        return f"{s} Sq. Ft."
    return s


def set_cell_text(cell, text, font_size=Pt(9), bold=False, color=RGBColor(0x33, 0x33, 0x33)):
    """Set cell text with consistent formatting."""
    cell.text = safe_str(text)
    for p in cell.text_frame.paragraphs:
        p.font.name = "Segoe UI"
        p.font.size = font_size
        p.font.bold = bold
        p.font.color.rgb = color


class PPTGenerator:
    def __init__(self, template_name="options format.pptx"):
        self.base_dir = os.path.dirname(os.path.abspath(__file__))
        self.template_path = os.path.join(self.base_dir, "templates", template_name)
        if not os.path.exists(self.template_path):
            # Fallback check
            alt_path = os.path.join(self.base_dir, "..", "LLM", "templates", template_name)
            if os.path.exists(alt_path):
                self.template_path = alt_path

    @staticmethod
    def _summary_table(slide):
        """The options table on a summary slide, or None."""
        for shape in slide.shapes:
            if shape.has_table:
                header = shape.table.rows[0].cells[0].text.strip().upper()
                if header.startswith(("SL", "SR", "NO")):
                    return shape.table
        return None

    def summary_capacity(self, slide):
        """
        How many options one summary table holds, read from the template.

        Hardcoding sixteen meant a change to the template silently truncated
        the deck; reading it means the template stays the authority.
        """
        table = self._summary_table(slide)
        return (len(table.rows) - 1) if table else 0

    def populate_summary_slide(self, slide9, matched_records, start=1):
        """
        Fill one summary table. `start` is the option number of the first row,
        so a deck whose options run past one table keeps numbering continuously
        across the slides that carry them.
        """
        tbl = self._summary_table(slide9)
        if tbl is None:
            return
        for idx in range(len(tbl.rows) - 1):
            row_cells = tbl.rows[idx + 1].cells
            if idx < len(matched_records):
                rec = matched_records[idx]
                prop_name = safe_str(rec.get('building_name', rec.get('property_name')),
                                     f'Option {start + idx}')
                loc_name = safe_str(rec.get('micromarket_category', rec.get('address_location')),
                                    'Bangalore')
                set_cell_text(row_cells[0], str(start + idx), font_size=Pt(8), bold=True)
                set_cell_text(row_cells[1], prop_name, font_size=Pt(8), bold=True)
                if len(row_cells) > 2:
                    set_cell_text(row_cells[2], loc_name, font_size=Pt(8))
            else:
                # Unused rows are blanked, so a short final table does not show
                # the placeholder text the template ships with.
                set_cell_text(row_cells[0], "")
                set_cell_text(row_cells[1], "")
                if len(row_cells) > 2:
                    set_cell_text(row_cells[2], "")

    # The building photograph placeholder on the option template. It is the
    # largest picture on the slide; the others are the logo and the footer rule.
    PHOTO_MIN_WIDTH_IN = 3.0

    @staticmethod
    def _fetch(url, timeout=20):
        """Read an image URL into memory, or None if it cannot be had."""
        if not url or not str(url).lower().startswith(("http://", "https://")):
            return None
        try:
            with urllib.request.urlopen(url, timeout=timeout) as response:
                if response.status != 200:
                    return None
                return response.read()
        except Exception:
            # A missing photograph is not a reason to fail a whole deck; the
            # template placeholder simply stays where it is.
            return None

    def place_photo(self, slide, url):
        """
        Swap the template placeholder for this building photograph.

        python-pptx cannot replace the image behind an existing picture shape,
        so the placeholder is measured, deleted, and a new picture inserted at
        the same position and size. Doing it by geometry rather than by shape
        name survives a template where the shapes have been renamed.
        """
        data = self._fetch(url)
        if not data:
            return False

        target = None
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                if Emu(shape.width).inches >= self.PHOTO_MIN_WIDTH_IN:
                    if target is None or shape.width > target.width:
                        target = shape
        if target is None:
            return False

        left, top, width, height = target.left, target.top, target.width, target.height
        target._element.getparent().remove(target._element)
        try:
            slide.shapes.add_picture(io.BytesIO(data), left, top,
                                     width=width, height=height)
        except Exception:
            return False
        return True

    # ------------------------------------------------------------ location
    # Geometry for a location slide, on a 13.33 x 7.5in canvas: title across
    # the top, the map filling the left, a panel of detail down the right.
    MAP_TITLE = (0.35, 0.22, 12.6, 0.62)
    MAP_IMAGE = (0.35, 1.05, 7.75, 5.85)
    MAP_PANEL = (8.35, 1.05, 4.60, 5.85)

    def _blank_layout(self, prs):
        """The emptiest layout the template offers, for a slide it has none of."""
        layouts = list(prs.slide_masters[0].slide_layouts)
        for layout in layouts:
            if layout.name.strip().lower() == "blank":
                return layout
        return layouts[-1]

    def add_location_slide(self, prs, title, image_bytes, panel_lines,
                           subtitle=""):
        """
        Build a location slide: a map on the left, notes down the right.

        The template ships no such slide, so this one is composed rather than
        cloned. Returns the slide, or None when there is no map to show - a
        deck is still a deck without one, and an empty frame reads as a bug.
        """
        if not image_bytes:
            return None

        slide = prs.slides.add_slide(self._blank_layout(prs))
        for shape in list(slide.shapes):
            shape.element.getparent().remove(shape.element)

        left, top, width, height = self.MAP_TITLE
        box = slide.shapes.add_textbox(Inches(left), Inches(top),
                                       Inches(width), Inches(height))
        frame = box.text_frame
        frame.word_wrap = True
        frame.text = title
        head = frame.paragraphs[0]
        head.font.size = Pt(20)
        head.font.bold = True
        head.font.name = "Segoe UI"
        head.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
        if subtitle:
            line = frame.add_paragraph()
            line.text = subtitle
            line.font.size = Pt(11)
            line.font.name = "Segoe UI"
            line.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

        left, top, width, height = self.MAP_IMAGE
        try:
            slide.shapes.add_picture(io.BytesIO(image_bytes), Inches(left),
                                     Inches(top), width=Inches(width),
                                     height=Inches(height))
        except Exception:
            return None

        if panel_lines:
            left, top, width, height = self.MAP_PANEL
            panel = slide.shapes.add_textbox(Inches(left), Inches(top),
                                             Inches(width), Inches(height))
            frame = panel.text_frame
            frame.word_wrap = True
            first = True
            for entry in panel_lines:
                # A tuple is a heading and its detail; a bare string is a line.
                heading, detail = entry if isinstance(entry, tuple) else ("", entry)
                if heading:
                    para = frame.paragraphs[0] if first else frame.add_paragraph()
                    para.text = heading
                    para.font.size = Pt(11)
                    para.font.bold = True
                    para.font.name = "Segoe UI"
                    para.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
                    para.space_before = Pt(0 if first else 8)
                    first = False
                if detail:
                    para = frame.paragraphs[0] if first else frame.add_paragraph()
                    para.text = detail
                    para.font.size = Pt(10)
                    para.font.name = "Segoe UI"
                    para.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                    first = False
        return slide

    def populate_option_slide(self, slide, rec, opt_num):
        """Populates Slide 10 option template with detailed property information."""
        prop_name = safe_str(rec.get('building_name', rec.get('property_name')), f'Option {opt_num}')
        dev_name = safe_str(rec.get('developer_landlord', rec.get('developer_name')), 'Developer Group')

        # The building photograph. Until now the record carried the URL and
        # nothing ever read it, so every option slide shipped with the template
        # stock image no matter what the database held.
        self.place_photo(slide, rec.get('_image_url') or rec.get('building_perspective_photo'))
        loc_name = safe_str(rec.get('address_location', rec.get('micromarket_category')), 'Bangalore')
        mm_name = safe_str(rec.get('micromarket_category'), 'Bangalore')
        
        # 1. Update Header Title — find the shape with "Option" text
        for shape in slide.shapes:
            if shape.has_text_frame:
                txt = shape.text_frame.text.strip()
                if "Option" in txt and ("–" in txt or "-" in txt or txt.startswith("Option")):
                    shape.text_frame.text = f"Option – {opt_num} – {prop_name}"
                    for p in shape.text_frame.paragraphs:
                        p.font.name = "Segoe UI"
                        p.font.size = Pt(20)
                        p.font.bold = True
                        p.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
                    break

        # 2. Update Tables
        for shape in slide.shapes:
            if shape.has_table:
                tbl = shape.table
                t_first_cell = tbl.rows[0].cells[0].text.strip().upper()
                
                # Table: Developer Info (3 rows x 2 cols)
                if 'NAME OF THE DEVELOPER' in t_first_cell or 'DEVELOPER' in t_first_cell:
                    set_cell_text(tbl.rows[0].cells[1], dev_name, bold=True)
                    set_cell_text(tbl.rows[1].cells[1], safe_str(rec.get('building_structure'), 'Grade A Building'))
                    set_cell_text(tbl.rows[2].cells[1], f"{loc_name}, {mm_name}" if loc_name != mm_name else loc_name)

                # Table: Building Details (4 rows x 2 cols)
                elif 'PROPOSED SPACE' in t_first_cell or 'BUILDING DETAIL' in t_first_cell:
                    if len(tbl.rows) >= 4:
                        set_cell_text(tbl.rows[1].cells[1], safe_str(rec.get('building_structure'), 'G + Upper Floors'))
                        set_cell_text(tbl.rows[2].cells[1], safe_str(rec.get('average_floor_plate_sqft'), 'Flexible floor plates'))
                        set_cell_text(tbl.rows[3].cells[1], safe_str(rec.get('floor_plate_efficiency'), '75-80%'))

                # Table: Commercial Details (9 rows x 2 cols)
                elif 'COMMERCIAL' in t_first_cell:
                    rent = format_currency(rec.get('quoted_rental_sqft_pm'))
                    cam = safe_str(rec.get('cam_charges_sqft_pm'), 'At Actuals')
                    if cam and 'INR' not in cam.upper() and cam not in ('At Actuals', 'Included in rent', 'TBD'):
                        try:
                            cam_num = float(cam.replace(',', ''))
                            cam = f"INR {cam_num:,.0f} / Sq. Ft. / Month"
                        except (ValueError, TypeError):
                            pass
                    parking_chg = safe_str(rec.get('car_parking_charges'), 'Included')
                    if parking_chg and 'INR' not in parking_chg.upper() and parking_chg != 'Included':
                        try:
                            pk_num = float(parking_chg.replace(',', '').replace('INR', '').strip())
                            parking_chg = f"INR {pk_num:,.0f} / Slot / Month"
                        except (ValueError, TypeError):
                            pass
                    
                    esc = format_percent(rec.get('rental_escalation'), '15% Every 36 Months')
                    sec_dep = format_months(rec.get('security_deposit_months'), '6 Months')
                    lease_term = format_months(rec.get('lease_tenure_months'), '60 Months')
                    lock_in = format_months(rec.get('lock_in_period_months'), '36 Months')
                    notice = format_months(rec.get('notice_period_months'), '6 Months')

                    rows_data = [
                        (1, rent, True),
                        (2, cam, False),
                        (3, parking_chg, False),
                        (4, esc, False),
                        (5, sec_dep, False),
                        (6, lease_term, False),
                        (7, lock_in, False),
                        (8, notice, False),
                    ]
                    for row_idx, val, is_bold in rows_data:
                        if row_idx < len(tbl.rows):
                            set_cell_text(tbl.rows[row_idx].cells[1], val, bold=is_bold)

                # Table: Details of Space Offered (8 rows x 2 cols)
                elif 'DETAILS OF THE SPACE' in t_first_cell or 'SPACE OFFERED' in t_first_cell:
                    area = format_quantity(rec.get('available_inventory_sqft'),
                                          rec.get('offered_seats'))
                    floor = safe_str(rec.get('floor_offered'), 'Multiple Floors')
                    fitout = safe_str(rec.get('fitout_details', rec.get('status')), 'Furnished')
                    power = safe_str(rec.get('power_kva'), '1 KVA / 100 Sq. Ft.')
                    power_bk = safe_str(rec.get('power_backup'), '100% DG Back-up')
                    parking_rat = safe_str(rec.get('car_parking_ratio'), '1:1000 Sq. Ft.')
                    timeline = safe_str(rec.get('timeline', rec.get('status')), 'Immediate')

                    rows_data = [
                        (1, area, True),
                        (2, floor, False),
                        (3, fitout, False),
                        (4, power, False),
                        (5, power_bk, False),
                        (6, parking_rat, False),
                        (7, timeline, False),
                    ]
                    for row_idx, val, is_bold in rows_data:
                        if row_idx < len(tbl.rows):
                            set_cell_text(tbl.rows[row_idx].cells[1], val, bold=is_bold)

        # 3. Update Highlights Box
        for shape in slide.shapes:
            if shape.has_text_frame:
                txt_lower = shape.text_frame.text.lower()
                if any(kw in txt_lower for kw in ["approvals", "distance", "eateries", "highlights", "nearby"]):
                    tf = shape.text_frame
                    tf.text = ""
                    
                    highlights = [
                        f"• Occupancy Certificate: {safe_str(rec.get('occupancy_certificate'), 'Yes')}",
                        f"• Address: {safe_str(rec.get('address_location'), loc_name)}",
                        f"• Micromarket: {mm_name}",
                        f"• Developer: {dev_name}",
                    ]
                    
                    contact = safe_str(rec.get('contact_person'))
                    if contact:
                        highlights.append(f"• Contact: {contact}")
                    email = safe_str(rec.get('official_email'))
                    if email:
                        highlights.append(f"• Email: {email}")
                    phone = safe_str(rec.get('contact_number'))
                    if phone:
                        highlights.append(f"• Phone: {phone}")
                    
                    # The workbook stores a lat/lng pair here. Rendered raw it
                    # read "Location: 12.981772994423237" in a client proposal,
                    # which says nothing the address bullet has not already said.
                    loc_map = safe_str(rec.get('location_map'))
                    is_coordinates = bool(re.fullmatch(
                        r"\s*-?\d+\.\d+\s*,\s*-?\d+\.\d+\s*", loc_map))
                    if loc_map and not is_coordinates and 'maps.google' not in loc_map:
                        highlights.append(f"• Location: {loc_map}")
                    
                    for idx, h_text in enumerate(highlights):
                        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                        p.text = h_text
                        p.font.name = "Segoe UI"
                        p.font.size = Pt(10)
                        p.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                        p.space_after = Pt(4)
                    break

    def _deep_clone_slide(self, prs, template_slide):
        """
        Deep-clones a slide, preserving shapes, tables, images and formatting.

        Two things here are easy to get wrong and both silently ruin the deck:

        1. The background must be copied as the `<p:bg>` child only. python-pptx's
           `slide.background._element` is the whole `<p:cSld>`, which *contains*
           the shape tree - replacing it swaps the freshly cloned shapes back for
           the template's, so every option slide ends up showing option 1.
        2. Relationships must be re-pointed. The copied `<a:blip r:embed="rId3">`
           still names rId3, which means nothing in the new part. python-pptx
           assigns rIds itself and offers no way to force one, so each carried
           relationship is added, the id it comes back with is recorded, and the
           copied XML is rewritten to match. The previous attempt passed the old
           rId as a third positional argument to `_add_relationship`, where the
           signature actually takes `is_external` - so every call failed, the
           bare `except` swallowed it, and no image relationship was ever
           carried. It went unnoticed because a slide whose photograph is
           replaced loses the broken reference anyway; only the options with no
           photograph of their own kept it, and those were the slides that
           would not open.
        """
        layout = template_slide.slide_layout
        new_slide = prs.slides.add_slide(layout)

        # Drop the layout's placeholder shapes; the clone brings its own.
        for sp in list(new_slide.shapes):
            sp.element.getparent().remove(sp.element)

        for shape_el in template_slide.shapes._spTree:
            tag = etree.QName(shape_el.tag).localname if isinstance(shape_el.tag, str) else ''
            if tag in ('sp', 'pic', 'graphicFrame', 'grpSp', 'cxnSp'):
                new_slide.shapes._spTree.append(copy.deepcopy(shape_el))

        # Carry the relationships over and remember what each was renamed to.
        remap = {}
        for rId, rel in template_slide.part.rels.items():
            if rel.reltype.endswith("slideLayout"):
                continue  # add_slide already wired the layout
            try:
                if rel.is_external:
                    new_id = new_slide.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
                else:
                    new_id = new_slide.part.rels.get_or_add(rel.reltype, rel.target_part)
                if new_id and new_id != rId:
                    remap[rId] = new_id
            except Exception:
                pass

        # Re-point every reference in the copied shapes at its new id. Any
        # attribute in the relationship namespace is rewritten, so pictures,
        # media, hyperlinks and charts all survive the copy.
        if remap:
            r_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
            for el in new_slide.shapes._spTree.iter():
                for attr, value in list(el.attrib.items()):
                    if attr.startswith("{%s}" % r_ns) and value in remap:
                        el.set(attr, remap[value])

        # Copy only the background element, never the containing cSld.
        try:
            ns = 'http://schemas.openxmlformats.org/presentationml/2006/main'
            src_cSld = template_slide._element.find('{%s}cSld' % ns)
            dst_cSld = new_slide._element.find('{%s}cSld' % ns)
            if src_cSld is not None and dst_cSld is not None:
                src_bg = src_cSld.find('{%s}bg' % ns)
                old_bg = dst_cSld.find('{%s}bg' % ns)
                if old_bg is not None:
                    dst_cSld.remove(old_bg)
                if src_bg is not None:
                    dst_cSld.insert(0, copy.deepcopy(src_bg))
        except Exception:
            pass

        return new_slide

    def generate_presentation(self, matched_records, summary_data=None,
                              client_name="Valued Client",
                              requirement_summary="Office Space Requirement",
                              output_filename="Commercial_Options_Deck.pptx",
                              overview_map=None, option_maps=None):
        """
        Loads 'options format.pptx', preserves slides 1-8 unchanged, populates summary slide 9,
        populates option slides from slide 10 onwards, and moves closing slides 11-13 to the end.

        `overview_map` is a prepared map of the whole shortlist; `option_maps`
        maps an option number to its own. Both are optional and independent,
        because which belongs in a proposal is the operator decision rather
        than a rule - options clustered in one market argue for the overview,
        a single option in an unfamiliar area for its own slide.
        """
        prs = pptx.Presentation(self.template_path)
        
        output_dir = os.path.join(self.base_dir, "generated_decks")
        os.makedirs(output_dir, exist_ok=True)
        output_path = os.path.join(output_dir, output_filename)

        if not matched_records:
            matched_records = [{
                'building_name': 'Sample Commercial Property',
                'developer_landlord': 'Prime Developer Group',
                'micromarket_category': 'Bangalore',
                'available_inventory_sqft': '15,000',
                'quoted_rental_sqft_pm': '100'
            }]

        id_list = prs.slides._sldIdLst
        snapshot = list(id_list)
        intro_els = snapshot[:8]
        closing_els = snapshot[10:13] if len(snapshot) > 12 else []

        def clone_of(source):
            """Clone a slide and hand back both it and its id element."""
            made = self._deep_clone_slide(prs, source)
            return made, list(id_list)[-1]

        # ---- summary: as many slides as the options need ---------------------
        # Every option appears in the summary. One table holds a fixed number of
        # rows, so the table is repeated rather than the list being cut to fit:
        # capping it at one slide hid options that were in the deck anyway, a few
        # slides further on.
        slide9 = prs.slides[8]
        capacity = self.summary_capacity(slide9) or len(matched_records)
        chunks = [matched_records[i:i + capacity]
                  for i in range(0, len(matched_records), capacity)] or [[]]

        # Clone every slide FIRST, from the untouched template, and only then
        # fill them. Populating the template before cloning it corrupted the
        # copies: swapping in a photograph deletes the placeholder picture and
        # adds a new one under a fresh relationship id, and the clones carried
        # the new XML while referencing a relationship that was never theirs -
        # the saved deck then failed to open the image at all.
        summary_slides, summary_els = [slide9], [snapshot[8]]
        for _ in chunks[1:]:
            made, el = clone_of(slide9)
            summary_slides.append(made)
            summary_els.append(el)
        for n, (sl, chunk) in enumerate(zip(summary_slides, chunks)):
            self.populate_summary_slide(sl, chunk, start=n * capacity + 1)

        # ---- one detail slide per option, however many there are -------------
        slide10_template = prs.slides[9]
        option_slides, option_els = [slide10_template], [snapshot[9]]
        for _ in range(2, len(matched_records) + 1):
            made, el = clone_of(slide10_template)
            option_slides.append(made)
            option_els.append(el)
        for idx, (sl, rec) in enumerate(zip(option_slides, matched_records), start=1):
            self.populate_option_slide(sl, rec, idx)

        # ---- location slides -------------------------------------------------
        # The overview follows the summary, so a client sees where the options
        # are before reading them one by one. An option map follows its own
        # slide, which is the only place it means anything.
        overview_els = []
        if overview_map:
            made = self.add_location_slide(
                prs, overview_map.get("title") or "Location overview",
                overview_map.get("image"), overview_map.get("lines") or [],
                overview_map.get("subtitle") or "")
            if made is not None:
                overview_els.append(list(id_list)[-1])

        option_map_els = {}
        for index in range(1, len(matched_records) + 1):
            prepared = (option_maps or {}).get(index)
            if not prepared:
                continue
            made = self.add_location_slide(
                prs, prepared.get("title") or "Location",
                prepared.get("image"), prepared.get("lines") or [],
                prepared.get("subtitle") or "")
            if made is not None:
                option_map_els[index] = list(id_list)[-1]

        # ---- put the deck back in reading order ------------------------------
        # Cloning appends, so summaries and options are interleaved at the end
        # until they are reordered. Rebuilding the whole list is clearer than
        # moving elements one at a time, and it keeps the closing slides last.
        interleaved = []
        for position, element in enumerate(option_els, start=1):
            interleaved.append(element)
            if position in option_map_els:
                interleaved.append(option_map_els[position])

        desired = intro_els + summary_els + overview_els + interleaved + closing_els
        for el in list(id_list):
            id_list.remove(el)
        for el in desired:
            id_list.append(el)

        prs.save(output_path)
        print(f"   -> Generated presentation deck at: {output_path}")
        return output_path
