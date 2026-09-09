import pptx
import copy

def generate_deck(template_path, matched_records, output_path):
    prs = pptx.Presentation(template_path)
    
    # Check initial slides count
    total_initial = len(prs.slides)
    print(f"Initial slides: {total_initial}")
    
    # Source Slide 10 (index 9)
    slide10_template = prs.slides[9]
    
    # Store closing slide elements (Slides 11, 12, 13 -> indices 10, 11, 12)
    closing_sld_ids = list(prs.slides._sldIdLst)[10:13]
    
    # We will keep Slide 10 as Option 1, and for Option 2..N we clone Slide 10
    option_slides = [slide10_template]
    
    for idx in range(2, len(matched_records) + 1):
        layout = slide10_template.slide_layout
        new_slide = prs.slides.add_slide(layout)
        # Remove default layout shapes
        for sp in list(new_slide.shapes):
            sp.element.getparent().remove(sp.element)
        # Deep copy all shapes from slide 10
        for shape in slide10_template.shapes:
            new_el = copy.deepcopy(shape.element)
            new_slide.shapes._spTree.append(new_el)
        option_slides.append(new_slide)
        
    # Now reorder _sldIdLst so closing_sld_ids are at the very end
    # Move closing_sld_ids to the end of _sldIdLst
    for sld_id_el in closing_sld_ids:
        prs.slides._sldIdLst.remove(sld_id_el)
        prs.slides._sldIdLst.append(sld_id_el)
        
    print("Final slide count:", len(prs.slides))
    print("Final slide order titles/first texts:")
    for idx, slide in enumerate(prs.slides, 1):
        txts = [p.text.strip() for shape in slide.shapes if shape.has_text_frame for p in shape.text_frame.paragraphs if p.text.strip()]
        print(f"  Slide {idx}: {txts[0] if txts else 'No text'}")
        
    prs.save(output_path)
    print("Saved reordered deck to:", output_path)

# Test with 3 sample options
generate_deck(
    r"c:\Users\devil\Desktop\full automation\LLM\templates\options format.pptx",
    [{"name": "Opt 1"}, {"name": "Opt 2"}, {"name": "Opt 3"}],
    r"c:\Users\devil\Desktop\full automation\LLM\generated_decks\test_reorder_out.pptx"
)
