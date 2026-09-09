import pptx
import copy

def clone_slide(prs, source_slide):
    """
    Creates a duplicate of source_slide in prs using source_slide.slide_layout and deep copying shapes.
    """
    layout = source_slide.slide_layout
    new_slide = prs.slides.add_slide(layout)
    
    # Clear layout elements created on the new slide
    for sp in list(new_slide.shapes):
        sp.element.getparent().remove(sp.element)
        
    for shape in source_slide.shapes:
        new_el = copy.deepcopy(shape.element)
        new_slide.shapes._spTree.append(new_el)
        
    return new_slide

pptx_path = r"c:\Users\devil\Desktop\full automation\LLM\templates\options format.pptx"
prs = pptx.Presentation(pptx_path)
s10 = prs.slides[9]

print(f"Original slide count: {len(prs.slides)}")
cloned_slide = clone_slide(prs, s10)
print(f"Cloned slide created! New slide count: {len(prs.slides)}")
prs.save(r"c:\Users\devil\Desktop\full automation\LLM\generated_decks\test_clone_output.pptx")
print("Saved cleanly.")
