import pptx

pptx_path = r"c:\Users\devil\Desktop\full automation\LLM\templates\options format.pptx"
prs = pptx.Presentation(pptx_path)

print("=== SLIDE 9 SHAPES & TABLES ===")
slide9 = prs.slides[8] # 0-indexed
for idx, shape in enumerate(slide9.shapes):
    print(f"Shape {idx}: {shape.name}, type={shape.shape_type}")
    if shape.has_text_frame:
        print("  Text:", [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()])
    if shape.has_table:
        print("  Table shape:", shape.table.rows[0].cells[0].text, f"{len(shape.table.rows)} rows x {len(shape.table.columns)} cols")
        for r_idx, r in enumerate(shape.table.rows):
            print(f"    Row {r_idx}:", [c.text.replace('\n', ' ').strip() for c in r.cells])

print("\n=== SLIDE 10 SHAPES & TABLES ===")
slide10 = prs.slides[9] # 0-indexed
for idx, shape in enumerate(slide10.shapes):
    print(f"Shape {idx}: {shape.name}, type={shape.shape_type}, pos=(L:{shape.left}, T:{shape.top}, W:{shape.width}, H:{shape.height})")
    if shape.has_text_frame:
        print("  Text:", [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()])
    if shape.has_table:
        print(f"  Table shape ({len(shape.table.rows)} rows x {len(shape.table.columns)} cols):")
        for r_idx, r in enumerate(shape.table.rows):
            print(f"    Row {r_idx}:", [c.text.replace('\n', ' ').strip() for c in r.cells])
