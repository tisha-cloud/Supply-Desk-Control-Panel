import pptx
import os

pptx_path = r"c:\Users\devil\Desktop\full automation\LLM\templates\options format.pptx"
prs = pptx.Presentation(pptx_path)

print(f"Total Slides in Template: {len(prs.slides)}")
print("=" * 70)

for idx, slide in enumerate(prs.slides, 1):
    print(f"\n--- SLIDE {idx} ---")
    texts = []
    tables = []
    images = 0
    shapes_info = []
    
    for s_idx, shape in enumerate(slide.shapes):
        shape_type = shape.shape_type
        name = shape.name
        pos = f"L:{shape.left}, T:{shape.top}, W:{shape.width}, H:{shape.height}"
        shapes_info.append(f"  Shape {s_idx+1}: {name} ({shape_type}) [{pos}]")
        
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if p.text.strip():
                    texts.append(p.text.strip())
        if shape.has_table:
            tbl_data = []
            for r in shape.table.rows:
                row_cells = [c.text.replace('\n', ' ').strip() for c in r.cells]
                tbl_data.append(row_cells)
            tables.append(tbl_data)
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            images += 1

    print("Texts found:")
    for t in texts[:10]:
        print("  -", t[:100])
    if len(texts) > 10:
        print(f"  ... (+ {len(texts)-10} more text blocks)")
        
    print(f"Images count: {images}")
    if tables:
        print(f"Tables count: {len(tables)}")
        for t_idx, t in enumerate(tables):
            print(f"  Table {t_idx+1} shape ({len(t)} rows x {len(t[0]) if t else 0} cols):")
            for r in t[:5]:
                print("    ", r)
