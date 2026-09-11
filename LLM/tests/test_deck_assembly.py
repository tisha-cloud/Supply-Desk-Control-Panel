"""
How the deck is assembled: how many slides, in what order, carrying what.

These build a real .pptx from the real template, because every defect they
cover was invisible in the object model and only showed up in the saved file.
"""
import hashlib
import os
import sys
import tempfile

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from ppt_generator import PPTGenerator  # noqa: E402


def records(count):
    return [{"building_name": "Building %02d" % i,
             "micromarket_category": "KRM",
             "developer_landlord": "Operator %02d" % i,
             "available_inventory_sqft": "",
             "offered_seats": 200 + i,
             "quoted_rental_sqft_pm": "INR 9,000 / seat / month"}
            for i in range(1, count + 1)]


def build(count, name):
    generator = PPTGenerator(template_name="options format.pptx")
    out = generator.generate_presentation(
        matched_records=records(count), client_name="Test",
        output_filename=name)
    return generator, Presentation(out), out


def option_slides(prs):
    return [s for s in prs.slides
            if any(sh.has_text_frame and sh.text_frame.text.strip().startswith("Option")
                   for sh in s.shapes)]


class TestEveryOptionGetsASlide:
    """
    The generator used to cap at sixteen - `min(len(records), 16)` - which
    silently dropped options the operator had deliberately kept in the
    shortlist.
    """

    def test_thirty_options_produce_thirty_slides(self):
        _, prs, _ = build(30, "test_thirty.pptx")
        assert len(option_slides(prs)) == 30

    def test_numbering_runs_to_the_end(self):
        _, prs, _ = build(20, "test_twenty.pptx")
        titles = [s.shapes[0].text_frame.text for s in option_slides(prs)
                  if s.shapes[0].has_text_frame]
        assert any("20" in t for t in titles)


class TestSummarySpillsRatherThanTruncating:
    def test_every_option_appears_in_a_summary_table(self):
        generator, prs, _ = build(30, "test_summary_spill.pptx")
        listed = []
        for slide in prs.slides:
            table = generator._summary_table(slide)
            if table is None:
                continue
            listed += [r.cells[0].text.strip() for r in list(table.rows)[1:]
                       if r.cells[0].text.strip()]
        assert listed == [str(n) for n in range(1, 31)]

    def test_numbering_continues_across_summary_slides(self):
        generator, prs, _ = build(30, "test_summary_numbering.pptx")
        tables = [generator._summary_table(s) for s in prs.slides]
        tables = [t for t in tables if t is not None]
        assert len(tables) == 2
        first = [r.cells[0].text.strip() for r in list(tables[0].rows)[1:]]
        second = [r.cells[0].text.strip() for r in list(tables[1].rows)[1:] if r.cells[0].text.strip()]
        assert first[0] == "1"
        assert second[0] == str(len(first) + 1)

    def test_one_summary_slide_when_the_options_fit(self):
        generator, prs, _ = build(5, "test_summary_single.pptx")
        tables = [generator._summary_table(s) for s in prs.slides]
        assert len([t for t in tables if t is not None]) == 1

    def test_capacity_is_read_from_the_template(self):
        """
        Hardcoding sixteen meant a template change silently truncated the deck.
        """
        generator, prs, _ = build(2, "test_capacity.pptx")
        slide9 = prs.slides[8]
        assert generator.summary_capacity(slide9) > 0


class TestRelationshipsSurviveCloning:
    """
    Cloned slides carried `<a:blip r:embed="rId3">` while their own part had no
    rId3, so the saved deck could not open the picture. It went unnoticed
    because a slide whose photograph is replaced loses the broken reference
    anyway - only options with no photograph kept it.
    """

    def test_no_slide_references_a_relationship_it_does_not_have(self):
        import re
        import zipfile
        _, _, path = build(12, "test_rels.pptx")
        archive = zipfile.ZipFile(path)
        broken = []
        for name in archive.namelist():
            if not re.match(r"ppt/slides/slide\d+\.xml$", name):
                continue
            xml = archive.read(name).decode("utf-8", "replace")
            embeds = set(re.findall(r'r:embed="(rId\d+)"', xml))
            rels_name = name.replace("slides/", "slides/_rels/") + ".rels"
            rels = (archive.read(rels_name).decode("utf-8", "replace")
                    if rels_name in archive.namelist() else "")
            have = set(re.findall(r'Id="(rId\d+)"', rels))
            if embeds - have:
                broken.append((name, sorted(embeds - have)))
        assert broken == []

    def test_every_option_slide_still_has_its_picture(self):
        _, prs, _ = build(12, "test_pictures.pptx")
        for slide in option_slides(prs):
            pictures = [sh for sh in slide.shapes
                        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
            assert pictures, "an option slide lost its imagery entirely"
            for picture in pictures:
                picture.image.blob          # raises if the relationship is dead


class TestPhotographPlacement:
    def test_a_missing_url_leaves_the_template_image_alone(self):
        generator = PPTGenerator(template_name="options format.pptx")
        prs = Presentation(generator.template_path)
        slide = prs.slides[9]
        before = len([sh for sh in slide.shapes
                      if sh.shape_type == MSO_SHAPE_TYPE.PICTURE])
        assert generator.place_photo(slide, None) is False
        assert generator.place_photo(slide, "") is False
        assert generator.place_photo(slide, "not-a-url") is False
        after = len([sh for sh in slide.shapes
                     if sh.shape_type == MSO_SHAPE_TYPE.PICTURE])
        assert after == before

    def test_an_unreachable_url_does_not_break_the_deck(self):
        """A dead photograph must never cost the client the whole proposal."""
        generator = PPTGenerator(template_name="options format.pptx")
        prs = Presentation(generator.template_path)
        assert generator.place_photo(
            prs.slides[9], "https://127.0.0.1:9/does-not-exist.jpg") is False
