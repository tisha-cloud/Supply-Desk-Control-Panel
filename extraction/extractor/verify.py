"""
Cross-checks the extracted Master File against the manually-maintained
reference workbook. The reference is read-only here - it is never modified.

The reference is itself incomplete (Prestige is floor-complete, Bagmane
partial, and Sattva/Embassy/Skav/99 Tech Park are name-only placeholders), so
this is a reconciliation report, not a pass/fail gate.
"""
import re
from difflib import SequenceMatcher

import openpyxl

from . import schema

MATCH_THRESHOLD = 0.82

_NOISE = re.compile(
    r"\b(the|group|groups|pvt|private|ltd|limited|llp|corp|corporation|developers?|"
    r"properties|property|estates?|building|block|tower|towers|park|phase)\b")


def _key(text):
    low = re.sub(r"[^a-z0-9 ]+", " ", str(text or "").lower())
    low = _NOISE.sub(" ", low)
    return re.sub(r"\s+", " ", low).strip()


def _similar(a, b):
    if not a or not b:
        return 0.0
    if a == b:
        return 1.0
    if a in b or b in a:
        return 0.94
    return SequenceMatcher(None, a, b).ratio()


def load_reference(path):
    """Group the reference Master File into one entry per building."""
    wb = openpyxl.load_workbook(path, data_only=True)
    if "Master File" not in wb.sheetnames:
        wb.close()
        return {}
    ws = wb["Master File"]
    buildings = {}
    # The sheet is hand-maintained: a building's identity columns are filled once
    # and left blank (or merged) on its continuation rows, so carry them forward.
    carried = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not any(row):
            continue
        # A new building name ends the previous building's run: its address,
        # structure and total size must not leak onto the next one.
        name_cell = schema.clean_name(row[3]) if len(row) > 3 else ""
        if name_cell and name_cell != carried.get(3):
            for idx in (0, 4, 5, 6):
                carried.pop(idx, None)
        for idx in (0, 2, 3, 4, 5, 6):
            value = schema.clean_name(row[idx]) if idx < len(row) else ""
            if value:
                carried[idx] = value
        developer = carried.get(2, "")
        building = carried.get(3, "")
        if not building:
            continue
        key = (_key(developer), _key(building))
        entry = buildings.setdefault(key, {
            "developer": developer,
            "building": building,
            "rows": 0,
            "total_size": None,
            "available_sqft": 0.0,
            "occupied_rows": 0,
            "has_floor_data": False,
        })
        entry["rows"] += 1
        total = schema.to_number(carried.get(6))
        if total and not entry["total_size"]:
            entry["total_size"] = total
        area = schema.to_number(row[8])
        timeline = str(row[10] or "").strip().lower()
        if row[7] or area:
            entry["has_floor_data"] = True
        if area:
            if timeline == "occupied":
                entry["occupied_rows"] += 1
            else:
                entry["available_sqft"] += area
    wb.close()
    return buildings


def _aggregate_extracted(master_rows):
    out = {}
    for row in master_rows:
        developer = row.get("_developer", "")
        building = row.get("_building", "")
        if not building:
            continue
        key = (_key(developer), _key(building))
        entry = out.setdefault(key, {
            "developer": developer,
            "building": building,
            "rows": 0,
            "total_size": row.get("_total_sqft"),
            "available_sqft": 0.0,
            "occupied_sqft": 0.0,
        })
        entry["rows"] += 1
        if not entry["total_size"] and row.get("_total_sqft"):
            entry["total_size"] = row.get("_total_sqft")
        area = row.get("_area_sqft") or 0
        if row.get("_occupancy") == "occupied":
            entry["occupied_sqft"] += area
        else:
            entry["available_sqft"] += area
    return out


def _best_match(ref_key, ref_entry, extracted):
    best, best_score = None, 0.0
    for key, entry in extracted.items():
        score = _similar(ref_key[1], key[1])
        if score < MATCH_THRESHOLD:
            continue
        if _similar(ref_key[0], key[0]) > 0.6:
            score += 0.05
        if score > best_score:
            best, best_score = key, score
    return best, best_score


def compare(master_rows, reference_path):
    """Returns verification rows for the audit sheet."""
    try:
        reference = load_reference(reference_path)
    except Exception as exc:
        return [{"Developer": "", "Building Name": "", "Status": "REFERENCE UNREADABLE",
                 "Note": str(exc)}]

    extracted = _aggregate_extracted(master_rows)
    results = []
    matched_keys = set()

    for ref_key, ref in sorted(reference.items()):
        key, score = _best_match(ref_key, ref, extracted)
        if not key:
            results.append({
                "Developer": ref["developer"],
                "Building Name": ref["building"],
                "Status": "IN REFERENCE ONLY",
                "Reference Rows": ref["rows"],
                "Extracted Rows": 0,
                "Reference Total Size": schema.sft(ref["total_size"]) if ref["total_size"] else "",
                "Extracted Total Size": "",
                "Reference Available (sft)": schema.sft(ref["available_sqft"]) if ref["available_sqft"] else "",
                "Extracted Available (sft)": "",
                "Note": ("no matching building in the source documents - the reference lists it "
                         "but no circulated file covers it"),
            })
            continue

        matched_keys.add(key)
        got = extracted[key]
        notes = []
        ref_total, got_total = ref["total_size"], got["total_size"]
        if ref_total and got_total and abs(ref_total - got_total) / max(ref_total, 1) > 0.05:
            notes.append("total size differs by %.0f%%" % (abs(ref_total - got_total) / ref_total * 100))
        elif ref_total and not got_total:
            notes.append("reference has a total size, extraction did not find one")
        elif got_total and not ref_total:
            notes.append("extraction recovered a total size the reference lacks")

        if ref["available_sqft"] and got["available_sqft"]:
            delta = abs(ref["available_sqft"] - got["available_sqft"])
            if delta / max(ref["available_sqft"], 1) > 0.10:
                notes.append("available area differs by %s" % schema.sft(delta))
        if not ref["has_floor_data"]:
            notes.append("reference row was a name-only placeholder; extraction filled it in")
        if score < 0.9:
            notes.append("fuzzy name match (%.2f)" % score)

        results.append({
            "Developer": ref["developer"],
            "Building Name": ref["building"],
            "Status": "MATCHED" if not notes else "MATCHED - REVIEW",
            "Reference Rows": ref["rows"],
            "Extracted Rows": got["rows"],
            "Reference Total Size": schema.sft(ref_total) if ref_total else "",
            "Extracted Total Size": schema.sft(got_total) if got_total else "",
            "Reference Available (sft)": schema.sft(ref["available_sqft"]) if ref["available_sqft"] else "",
            "Extracted Available (sft)": schema.sft(got["available_sqft"]) if got["available_sqft"] else "",
            "Note": "; ".join(notes),
        })

    for key, got in sorted(extracted.items()):
        if key in matched_keys:
            continue
        results.append({
            "Developer": got["developer"],
            "Building Name": got["building"],
            "Status": "NEW FROM EXTRACTION",
            "Reference Rows": 0,
            "Extracted Rows": got["rows"],
            "Reference Total Size": "",
            "Extracted Total Size": schema.sft(got["total_size"]) if got["total_size"] else "",
            "Reference Available (sft)": "",
            "Extracted Available (sft)": schema.sft(got["available_sqft"]) if got["available_sqft"] else "",
            "Note": "not present in the manual reference sheet",
        })

    order = {"IN REFERENCE ONLY": 0, "MATCHED - REVIEW": 1, "MATCHED": 2, "NEW FROM EXTRACTION": 3}
    results.sort(key=lambda r: (order.get(r["Status"], 9), r["Developer"], r["Building Name"]))
    return results
