import pptx
import re
import os

def parse_pptx_file(item):
    """
    Parses PPTX presentation decks using python-pptx.
    """
    abs_path = item['abs_path']
    rel_path = item['rel_path']
    developer_name = item['developer_name']
    period = item['period']
    folder_category = item['folder_category']
    
    records = []
    
    try:
        prs = pptx.Presentation(abs_path)
    except Exception as e:
        return [{
            'developer_name': developer_name,
            'folder_category': folder_category,
            'property_name': 'Error reading PPTX',
            'raw_remarks': f'Failed to open PPTX: {str(e)}',
            'source_file': rel_path,
            'source_format': 'PPTX',
            'extraction_method': 'PPTX Error',
            'report_period': period
        }]

    slides_data = []
    
    for slide_idx, slide in enumerate(prs.slides):
        slide_text_list = []
        slide_tables = []
        
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    txt = paragraph.text.strip()
                    if txt:
                        slide_text_list.append(txt)
            if shape.has_table:
                tbl = shape.table
                tbl_rows = []
                for row in tbl.rows:
                    cells = [c.text.replace('\n', ' ').strip() for c in row.cells]
                    if any(cells):
                        tbl_rows.append(cells)
                if tbl_rows:
                    slide_tables.append(tbl_rows)
                    
        full_slide_text = " | ".join(slide_text_list)
        slides_data.append({
            'slide_num': slide_idx + 1,
            'text': full_slide_text,
            'tables': slide_tables
        })

    # Global deck properties dictionary
    deck_property = {
        'developer_name': developer_name,
        'folder_category': folder_category,
        'property_name': os.path.splitext(item['filename'])[0],
        'source_file': rel_path,
        'source_format': 'PPTX',
        'extraction_method': 'PPTX Deck Analysis',
        'report_period': period
    }
    
    all_deck_text = " ".join([s['text'] for s in slides_data])
    
    # Key extraction regex pattern matching
    loc_m = re.search(r'(?:Location|Situated in)[\s|\:]+([^\.\|]+)', all_deck_text, re.I)
    area_m = re.search(r'([\d,]+)\s*(?:Sq\.?\s*Ft|sqft|sft)', all_deck_text, re.I)
    rent_m = re.search(r'(?:Rent|Rate)[\s|\:\-]*Rs\.?\s*([\d,]+(?:\.\d+)?)', all_deck_text, re.I)
    floor_m = re.search(r'(\d+(?:st|nd|rd|th)?\s*floor|Ground\s*floor|GF)', all_deck_text, re.I)
    fitout_m = re.search(r'(Fully\s*furnished|Plug\s*and\s*Play|Warm\s*shell|Bare\s*shell|Managed)', all_deck_text, re.I)
    parking_m = re.search(r'(?:Car\s*Parking|Parking)[\s|\:\-]*([^\.\|]+)', all_deck_text, re.I)
    contact_m = re.search(r'([A-Z][a-z]+\s+[A-Z][a-z]+)?\s*[\|\-]?\s*(?:Cell|Mobile|Ph|T)[\s:\-]+([\d\s\+\-]+)', all_deck_text, re.I)
    
    if loc_m: deck_property['micromarket_location'] = loc_m.group(1).strip()
    if area_m: deck_property['available_area_sqft'] = area_m.group(1).strip()
    if rent_m: deck_property['rent_rate_sqft_pm'] = rent_m.group(1).strip()
    if floor_m: deck_property['floor_details'] = floor_m.group(1).strip()
    if fitout_m: deck_property['fitout_status'] = fitout_m.group(1).strip()
    if parking_m: deck_property['car_parking_details'] = parking_m.group(1).strip()
    if contact_m: deck_property['contact_details'] = f"{contact_m.group(1) or ''} {contact_m.group(2)}".strip()
    
    # Process tables in slides
    table_records = []
    for s in slides_data:
        for tbl in s['tables']:
            if len(tbl) >= 2:
                headers = tbl[0]
                for r in tbl[1:]:
                    rec = dict(deck_property)
                    rec['extraction_method'] = f"PPTX Table (Slide {s['slide_num']})"
                    pairs = []
                    for c_idx, cell in enumerate(r):
                        if c_idx < len(headers) and cell:
                            h = headers[c_idx]
                            pairs.append(f"{h}: {cell}")
                            h_u = h.upper()
                            if 'LOCATION' in h_u: rec['micromarket_location'] = cell
                            elif 'AREA' in h_u or 'SIZE' in h_u: rec['available_area_sqft'] = cell
                            elif 'RENT' in h_u or 'COMMERCIAL' in h_u: rec['rent_rate_sqft_pm'] = cell
                            elif 'FLOOR' in h_u: rec['floor_details'] = cell
                            elif 'STATUS' in h_u or 'DESCRIPTION' in h_u: rec['fitout_status'] = cell
                    if pairs:
                        rec['raw_remarks'] = " | ".join(pairs)
                        table_records.append(rec)

    if table_records:
        return table_records
        
    deck_property['raw_remarks'] = all_deck_text[:500]
    return [deck_property]

if __name__ == '__main__':
    import glob
    files = glob.glob(r'c:\Users\devil\Desktop\Data-extract\BLR_Builders_Developers_Supply\**\*.pptx', recursive=True)
    for f in files:
        item = {
            'abs_path': f,
            'rel_path': os.path.relpath(f, r'c:\Users\devil\Desktop\Data-extract\BLR_Builders_Developers_Supply'),
            'developer_name': 'Indraprastha Shelters',
            'folder_category': 'Indraprastha Shelters Pvt. Ltd',
            'filename': os.path.basename(f),
            'period': 'July 2026'
        }
        res = parse_pptx_file(item)
        print(f"PPTX file {item['filename']} -> Extracted {len(res)} records.")
        print("  Sample:", res[0])
