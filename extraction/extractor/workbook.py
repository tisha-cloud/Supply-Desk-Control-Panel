"""
Writes the output workbook in the shape of
'Blr - Conventional - Options Sheet -2026.xlsx'.

Sheets produced:
    Builder - Developer - PoC   contacts harvested from the source documents
    Master File                 the deliverable: Bangalore conventional office, one row per floor
    Commercial - Sale           assets explicitly marketed for sale
    Other Assets                retail / industrial / warehouse / out-of-Bangalore stock
    Managed Office Space (x4)   blank comparison templates carried over from the reference
    Extraction Audit            provenance for every Master File row
    Occupancy Derivation        how each balance row was computed, and what blocked it
    Verification                extracted vs the reference workbook
    File Coverage               every source file and what came out of it
"""
from copy import copy

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .schema import MASTER_COLUMNS

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
DERIVED_FILL = PatternFill("solid", fgColor="FFF2CC")
FLAG_FILL = PatternFill("solid", fgColor="FCE4E4")
BODY_FONT = Font(size=10)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MASTER_WIDTHS = [20, 9, 24, 34, 34, 18, 22, 22, 20, 20, 16]

POC_COLUMNS = ["Sl. No", "Builder / Developer", "Building Name / Location",
               "Contact Person", "Contact Number", "Email ID", "Inventory Status"]
SALE_COLUMNS = ["Sl. No", "Builder - Developer - Portfolio", "Building Name / Location",
                "Address / Location", "Contact Person", "Contact Number", "Email ID"]


def _style_header(ws, ncols, freeze=True, autofilter=True):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 30
    if freeze:
        ws.freeze_panes = "A2"
    if autofilter and ws.max_row > 1:
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(ncols), ws.max_row)


def _write_table(ws, columns, records, widths=None, wrap_cols=()):
    ws.append(list(columns))
    for record in records:
        ws.append([record.get(col, "") for col in columns])
    for idx, col in enumerate(columns, start=1):
        letter = get_column_letter(idx)
        if widths and idx <= len(widths):
            ws.column_dimensions[letter].width = widths[idx - 1]
        else:
            ws.column_dimensions[letter].width = max(14, min(46, len(str(col)) + 8))
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(columns)):
        for cell in row:
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center",
                                       wrap_text=cell.column_letter in wrap_cols)
    _style_header(ws, len(columns))


def write_master_sheet(ws, rows):
    ws.append(list(MASTER_COLUMNS))
    for row in rows:
        ws.append([row.get(col, "") for col in MASTER_COLUMNS])
        if row.get("_derived"):
            for col in range(1, len(MASTER_COLUMNS) + 1):
                ws.cell(row=ws.max_row, column=col).fill = DERIVED_FILL
    for idx, width in enumerate(MASTER_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(MASTER_COLUMNS)):
        for cell in row:
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center",
                                       wrap_text=cell.column_letter in ("D", "E"))
    _style_header(ws, len(MASTER_COLUMNS))


def copy_template_sheet(src_ws, dst_ws):
    """Carry a reference sheet across with its values, styling and geometry."""
    for row in src_ws.iter_rows():
        for cell in row:
            target = dst_ws.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                target.font = copy(cell.font)
                target.fill = copy(cell.fill)
                target.border = copy(cell.border)
                target.alignment = copy(cell.alignment)
                target.number_format = cell.number_format
    for key, dim in src_ws.column_dimensions.items():
        dst_ws.column_dimensions[key].width = dim.width
        dst_ws.column_dimensions[key].hidden = dim.hidden
    for key, dim in src_ws.row_dimensions.items():
        dst_ws.row_dimensions[key].height = dim.height
    for merged in list(src_ws.merged_cells.ranges):
        try:
            dst_ws.merge_cells(str(merged))
        except Exception:
            pass
    dst_ws.freeze_panes = src_ws.freeze_panes


def build_workbook(out_path, master_rows, poc_rows, sale_rows, other_rows,
                   audit_rows, derivation_rows, verification_rows, coverage_rows,
                   reference_path=None):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    _write_table(wb.create_sheet("Builder - Developer - PoC"), POC_COLUMNS, poc_rows,
                 widths=[8, 26, 40, 24, 18, 34, 16])

    write_master_sheet(wb.create_sheet("Master File"), master_rows)

    _write_table(wb.create_sheet("Commercial - Sale"), SALE_COLUMNS, sale_rows,
                 widths=[8, 26, 34, 30, 22, 18, 34])

    other_cols = ["Asset Type", "City", "Micro-Market", "Builder / Developer", "Building Name",
                  "Address / Location", "Building Structure", "Total Building Size [In Sq. Ft.]",
                  "Available  Floors", "Available - Area in Sft", "Condition", "Timeline",
                  "Source File"]
    _write_table(wb.create_sheet("Other Assets"), other_cols, other_rows,
                 widths=[14, 14, 18, 22, 32, 32, 16, 22, 20, 20, 18, 16, 46])

    # Blank comparison templates carried over untouched from the reference workbook.
    if reference_path:
        try:
            ref = openpyxl.load_workbook(reference_path, data_only=True)
            for name in ref.sheetnames:
                if name.startswith("Managed Office Space"):
                    copy_template_sheet(ref[name], wb.create_sheet(name))
            ref.close()
        except Exception as exc:
            wb.create_sheet("Managed Office Space").cell(
                row=1, column=1, value="Template copy failed: %s" % exc)

    audit_cols = ["Developer", "Building Name", "Micro-Market", "City", "Asset Type",
                  "Transaction", "Floor", "Area in Sft", "Condition", "Condition Detail",
                  "Timeline", "Occupancy", "Rent /sft/mo", "CAM /sft/mo",
                  "Total Building Size", "Size Basis", "Disclosure Mode", "Derivation",
                  "Evidence", "Notes", "Report Period", "Source File"]
    _write_table(wb.create_sheet("Extraction Audit"), audit_cols, audit_rows,
                 widths=[20, 30, 16, 12, 12, 12, 16, 16, 16, 30, 14, 12, 12, 12,
                         18, 26, 16, 40, 28, 30, 14, 46])

    deriv_cols = ["Developer", "Building Name", "Total Building Size (sft)", "Total Size Source",
                  "Listed Available (sft)", "Listed Occupied (sft)", "Derived Balance (sft)",
                  "Balance Row Emitted", "Disclosure Mode", "Issue", "Source File"]
    ws_deriv = wb.create_sheet("Occupancy Derivation")
    _write_table(ws_deriv, deriv_cols, derivation_rows,
                 widths=[20, 32, 20, 16, 20, 20, 20, 18, 18, 60, 46])
    issue_idx = deriv_cols.index("Issue")
    source_idx = deriv_cols.index("Total Size Source")
    for row in ws_deriv.iter_rows(min_row=2, max_row=ws_deriv.max_row, max_col=len(deriv_cols)):
        if row[issue_idx].value:
            for cell in row:
                cell.fill = FLAG_FILL
        elif row[source_idx].value == "Estimated":
            for cell in row:
                cell.fill = DERIVED_FILL

    verify_cols = ["Developer", "Building Name", "Status", "Reference Rows", "Extracted Rows",
                   "Reference Total Size", "Extracted Total Size", "Reference Available (sft)",
                   "Extracted Available (sft)", "Note"]
    _write_table(wb.create_sheet("Verification"), verify_cols, verification_rows,
                 widths=[20, 34, 22, 16, 16, 20, 20, 22, 22, 52])

    coverage_cols = ["Developer", "Source File", "Type", "Size (MB)", "Report Period",
                     "Buildings", "Master Rows", "Other Rows", "Status", "Detail"]
    ws_cov = wb.create_sheet("File Coverage")
    _write_table(ws_cov, coverage_cols, coverage_rows,
                 widths=[22, 52, 10, 11, 14, 11, 12, 11, 18, 50])
    for row in ws_cov.iter_rows(min_row=2, max_row=ws_cov.max_row, max_col=len(coverage_cols)):
        if str(row[8].value).startswith(("FAILED", "NO DATA")):
            for cell in row:
                cell.fill = FLAG_FILL

    wb.save(out_path)
    return out_path
