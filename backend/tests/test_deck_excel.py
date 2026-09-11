"""
The spreadsheet export, and the formatting it shares with the deck.

The sheet deliberately reproduces the per-market layout of the supply workbook
- a fixed spine of particulars down column A, one column per option across -
so the output drops straight into the workbook the team already reads.
"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook  # noqa: E402

from services import deck_excel  # noqa: E402
from services.deck_service import _fmt_int, _fmt_sqft  # noqa: E402


class TestFormatters:
    def test_zero_area_is_blank_not_a_figure(self):
        """
        A seat-based option has no area. Printing the zero put
        "Area Offered: 0 Sq. Ft." on a suite that had 40 desks available.
        """
        assert _fmt_sqft(0) == ""
        assert _fmt_sqft(None) == ""
        assert _fmt_int(0) == ""

    def test_indian_digit_grouping(self):
        assert _fmt_sqft(100000) == "1,00,000 Sq. Ft."
        assert _fmt_sqft(30000) == "30,000 Sq. Ft."
        assert _fmt_int(7000) == "7,000"


def record(name, market, seats="", area="", landlord="An Operator"):
    return {"building_name": name, "micromarket_category": market,
            "developer_landlord": landlord, "address_location": "%s Road" % market,
            "available_inventory_sqft": area, "offered_seats": seats,
            "quoted_rental_sqft_pm": "INR 9,000 / seat / month",
            "status": "Fully Furnished", "timeline": "Immediate",
            "lease_tenure_months": "36 Months", "contact_person": "A Person"}


RECORDS = [
    record("Prestige Cube", "KRM", seats=167),
    record("Padmavathi Complex", "KRM", seats=171),
    record("Sunriver", "Indiranagar", seats=189),
    record("Tower A", "CBD", area="30,000 Sq. Ft."),
]
CRITERIA = {"title": "150 seats in Koramangala",
            "product_note": "150 seats is 15 or more, so this is a managed office."}


def build():
    path = os.path.join(tempfile.mkdtemp(), "options.xlsx")
    deck_excel.build(RECORDS, CRITERIA, "Acme Labs", "150 seats", path)
    return load_workbook(path)


class TestOneSheetPerMicroMarket:
    def test_a_sheet_for_each_market_plus_the_brief(self):
        wb = build()
        assert wb.sheetnames == ["Requirement", "KRM", "Indiranagar", "CBD"]

    def test_markets_appear_in_ranked_order(self):
        """
        The market holding the best-fitting option leads, because the records
        arrive already ranked and the sheets follow them.
        """
        assert build().sheetnames[1] == "KRM"

    def test_each_sheet_holds_only_its_own_options(self):
        wb = build()
        assert wb["KRM"].max_column == 3       # label column + two options
        assert wb["Indiranagar"].max_column == 2
        assert wb["CBD"].max_column == 2


class TestTheSpineMatchesTheSupplyWorkbook:
    def test_the_row_labels_are_the_workbook_labels(self):
        ws = build()["KRM"]
        labels = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
        assert labels[0] == "Particulars"
        assert labels[1] == "Building Name"
        assert labels[2] == "Address"
        assert labels[3] == "Building Perspective/ Photograph"
        assert "Building Details" in labels
        assert "Proposed Space" in labels
        assert "Commercial Terms" in labels
        assert "Contact Details" in labels
        assert labels[-1] == "Official Email ID"

    def test_the_spine_is_thirty_five_rows(self):
        assert build()["KRM"].max_row == 35

    def test_values_run_across_not_down(self):
        ws = build()["KRM"]
        assert ws.cell(row=2, column=2).value == "Prestige Cube"
        assert ws.cell(row=2, column=3).value == "Padmavathi Complex"

    def test_section_bands_are_red(self):
        ws = build()["KRM"]
        labels = {ws.cell(row=r, column=1).value: r for r in range(1, ws.max_row + 1)}
        for band in ("Particulars", "Building Details", "Commercial Terms"):
            cell = ws.cell(row=labels[band], column=1)
            assert cell.fill.start_color.rgb.endswith("FF0000")
            assert cell.font.bold


class TestOptionNumbering:
    def test_numbers_run_across_the_whole_proposal_not_per_sheet(self):
        """
        The deck calls the third option "Option 3". A sheet that restarted at 1
        for its own market would disagree with the deck about one building.
        """
        wb = build()
        assert wb["KRM"].cell(row=1, column=2).value == "Option 1"
        assert wb["KRM"].cell(row=1, column=3).value == "Option 2"
        assert wb["Indiranagar"].cell(row=1, column=2).value == "Option 3"
        assert wb["CBD"].cell(row=1, column=2).value == "Option 4"


class TestQuantity:
    def test_seats_when_the_option_is_sold_by_the_seat(self):
        ws = build()["KRM"]
        row = next(r for r in range(1, ws.max_row + 1)
                   if str(ws.cell(row=r, column=1).value).startswith("Offered Seats"))
        assert ws.cell(row=row, column=2).value == "167"

    def test_area_when_there_are_no_seats(self):
        ws = build()["CBD"]
        row = next(r for r in range(1, ws.max_row + 1)
                   if str(ws.cell(row=r, column=1).value).startswith("Offered Seats"))
        assert ws.cell(row=row, column=2).value == "30,000 Sq. Ft."


class TestPhotographs:
    def test_a_missing_or_broken_url_never_fails_the_export(self):
        assert deck_excel._fetch(None) is None
        assert deck_excel._fetch("") is None
        assert deck_excel._fetch("not-a-url") is None
        assert deck_excel._fetch("https://127.0.0.1:9/nope.jpg") is None

    def test_the_photo_row_is_tall_enough_to_show_one(self):
        ws = build()["KRM"]
        assert ws.row_dimensions[4].height == deck_excel.PHOTO_ROW_HEIGHT


class TestTheBrief:
    def test_it_records_who_and_why(self):
        ws = build()["Requirement"]
        text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
        assert "Acme Labs" in text
        assert "150 seats" in text
        assert "Options included" in text
        assert "KRM (2)" in text
