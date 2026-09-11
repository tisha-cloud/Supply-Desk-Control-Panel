"""
The shortlist as an options sheet, in the format the desk already uses.

This mirrors the per-market sheets of "BLR - Managed Office Space Supply" -
CBD, Indiranagar, KRM and the rest: a fixed spine of particulars down column A
and one column per option across, photograph included. Reproducing that layout
rather than inventing a grid means the output drops straight into the workbook
the team already reads, and a reviewer does not have to learn a second shape
for the same information.

Both this and the PowerPoint are rendered from the same option records, so a
client sent both cannot find them disagreeing.
"""
import io
import os
import urllib.request
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Taken from the source workbook rather than chosen: the sheets these land
# beside use exactly this red, this face and this size.
BAND_FILL = PatternFill("solid", fgColor="FF0000")
BAND_FONT = Font(name="Century Gothic", size=8, bold=True, color="FFFFFF")
LABEL_FONT = Font(name="Century Gothic", size=8, bold=True)
VALUE_FONT = Font(name="Calibri", size=8)
TITLE_FONT = Font(name="Century Gothic", size=14, bold=True)
THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

LABEL_WIDTH = 40.3
OPTION_WIDTH = 37.1
PHOTO_ROW_HEIGHT = 141.8

# The spine, in order. (row label, record key). A key of None is a section
# band; the photograph row is handled separately because it holds an image
# rather than a value.
PHOTO_ROW = "Building Perspective/ Photograph"
SPINE: List[Tuple[str, Optional[str]]] = [
    ("Particulars", "_option_label"),
    ("Building Name", "building_name"),
    ("Address", "address_location"),
    (PHOTO_ROW, None),
    ("Building Details", None),
    ("Developer/ Landlord", "developer_landlord"),
    ("Building Structure", "building_structure"),
    ("Total Building Size (SBA) [In Sq. Ft.]", "total_building_size_sqft"),
    ("Average Floor Plate Size [In Sq. Ft.]", "average_floor_plate_sqft"),
    ("Floor Plate Efficiency [Approx.]", "floor_plate_efficiency"),
    ("Power [In KVA]", "power_kva"),
    ("Power Back-up", "power_backup"),
    ("Proposed Space", None),
    ("Offered Seats / Area Offered [In Sq. Ft.]", "_quantity"),
    ("Floor Offered", "floor_offered"),
    ("Status", "status"),
    ("Fit-out Details", "fitout_details"),
    ("Timeline", "timeline"),
    ("Car Parking Ratio", "car_parking_ratio"),
    ("Commercial Terms", None),
    ("Quoted Rental [INR/ Sq. Ft./ Month]", "quoted_rental_sqft_pm"),
    ("CAM Charges [INR/ Sq. Ft./ Month]", "cam_charges_sqft_pm"),
    ("Car Parking Charges [INR/ Slot/ Month]", "car_parking_charges"),
    ("Rental Escalation [% In Months]", "rental_escalation"),
    ("Interest Free Refundable Security Deposit [In Months]", "security_deposit_months"),
    ("Lease Tenure [In Months]", "lease_tenure_months"),
    ("Lock-In Period [In Months]", "lock_in_period_months"),
    ("Notice Period for Termination [In Months]", "notice_period_months"),
    ("Occupancy Certificate Available (Yes / No)", "occupancy_certificate"),
    ("Location Map", "location_map"),
    ("", None),
    ("Contact Details", None),
    ("Contact Person", "contact_person"),
    ("Contact Number", "contact_number"),
    ("Official Email ID", "official_email"),
]

# Rows that carry the red band across the whole sheet. The first three are the
# header block; the rest are section dividers.
BANDED = {"Particulars", "Building Name", "Address",
          "Building Details", "Proposed Space", "Commercial Terms",
          "Contact Details"}

# A photograph is scaled to sit inside its cell: roughly 265 x 189 pixels at
# the column width and row height above.
PHOTO_BOX = (255, 180)


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def _fetch(url: str, timeout: int = 20) -> Optional[bytes]:
    """Read a photograph, or None. A missing image never fails the export."""
    if not url or not str(url).lower().startswith(("http://", "https://")):
        return None
    try:
        with urllib.request.urlopen(url, timeout=timeout) as response:
            return response.read() if response.status == 200 else None
    except Exception:
        return None


def _scaled(data: bytes) -> Optional[XLImage]:
    """An openpyxl image scaled to fit the photograph cell, aspect preserved."""
    try:
        image = XLImage(io.BytesIO(data))
    except Exception:
        return None
    box_w, box_h = PHOTO_BOX
    try:
        ratio = min(box_w / float(image.width), box_h / float(image.height), 1.0)
        image.width = int(image.width * ratio)
        image.height = int(image.height * ratio)
    except Exception:
        image.width, image.height = box_w, box_h
    return image


def _market_of(record: Dict[str, Any]) -> str:
    return (record.get("micromarket_category") or "Other").strip() or "Other"


def _quantity(record: Dict[str, Any]) -> str:
    """
    What is on offer, seats or area.

    The source sheets put both in one row, whichever the operator quotes, so a
    seat-based centre reads "115" and a conventional floor "30,000 Sq. Ft."
    """
    seats = record.get("offered_seats")
    if seats:
        return _text(seats)
    return _text(record.get("available_inventory_sqft"))


def build(records: List[Dict[str, Any]],
          criteria: Dict[str, Any],
          client_name: str,
          query: str,
          out_path: str) -> str:
    """Write the options workbook and return the path it was written to."""
    wb = Workbook()
    _write_brief(wb.active, records, criteria, client_name, query)
    wb.active.title = "Requirement"

    # One sheet per micro-market, in the order the options were ranked, so the
    # market carrying the best fit comes first.
    order: List[str] = []
    grouped: Dict[str, List[Dict[str, Any]]] = {}
    for record in records:
        market = _market_of(record)
        if market not in grouped:
            grouped[market] = []
            order.append(market)
        grouped[market].append(record)

    # Option numbers run across the whole proposal, not per sheet. The deck
    # calls the thirteenth option "Option 13"; a sheet that restarted at 1 for
    # its own market would disagree with the deck about the same building.
    for position, record in enumerate(records, start=1):
        record["_option_number"] = position

    for market in order:
        _write_market_sheet(wb.create_sheet(_sheet_name(market, wb)), grouped[market])

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
    return out_path


def _sheet_name(market: str, wb: Workbook) -> str:
    """Excel forbids several characters in a sheet name and caps it at 31."""
    name = "".join(c for c in market if c not in "[]:*?/\\")[:31] or "Options"
    if name in wb.sheetnames:
        suffix = 2
        while ("%s %d" % (name[:28], suffix)) in wb.sheetnames:
            suffix += 1
        name = "%s %d" % (name[:28], suffix)
    return name


def _write_market_sheet(ws, records: List[Dict[str, Any]]) -> None:
    ws.column_dimensions["A"].width = LABEL_WIDTH
    for index in range(len(records)):
        ws.column_dimensions[get_column_letter(index + 2)].width = OPTION_WIDTH

    width = len(records) + 1

    for row_index, (label, key) in enumerate(SPINE, start=1):
        banded = label in BANDED
        cell = ws.cell(row=row_index, column=1, value=label)
        cell.font = BAND_FONT if banded else LABEL_FONT
        cell.border = BORDER
        cell.alignment = Alignment(
            horizontal="center" if banded else "left",
            vertical="center", wrap_text=True)
        if banded:
            cell.fill = BAND_FILL

        for offset, record in enumerate(records):
            column = offset + 2
            value = ""
            if key == "_option_label":
                value = "Option %d" % (record.get("_option_number") or offset + 1)
            elif key == "_quantity":
                value = _quantity(record)
            elif key:
                value = _text(record.get(key))

            target = ws.cell(row=row_index, column=column, value=value)
            target.border = BORDER
            if banded:
                target.font = BAND_FONT
                target.fill = BAND_FILL
                target.alignment = Alignment(horizontal="center", vertical="center",
                                             wrap_text=True)
            else:
                target.font = VALUE_FONT
                target.alignment = Alignment(horizontal="center", vertical="center",
                                             wrap_text=True)

        if label == PHOTO_ROW:
            ws.row_dimensions[row_index].height = PHOTO_ROW_HEIGHT
            for offset, record in enumerate(records):
                data = _fetch(record.get("_image_url")
                              or record.get("building_perspective_photo"))
                image = _scaled(data) if data else None
                if image is not None:
                    ws.add_image(image, "%s%d" % (get_column_letter(offset + 2), row_index))

    ws.freeze_panes = "B1"
    # Nothing below the spine belongs to the sheet, so the print area stops
    # where the options do.
    ws.print_area = "A1:%s%d" % (get_column_letter(width), len(SPINE))


def _write_brief(ws, records, criteria, client_name, query) -> None:
    """Who it is for and what was asked, ahead of the market sheets."""
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 78

    ws["A1"] = criteria.get("title") or "Office Space Options"
    ws["A1"].font = TITLE_FONT
    ws.append([])

    def line(label: str, value: Any) -> None:
        ws.append([label, _text(value)])
        ws.cell(row=ws.max_row, column=1).font = LABEL_FONT
        ws.cell(row=ws.max_row, column=2).font = VALUE_FONT
        ws.cell(row=ws.max_row, column=2).alignment = Alignment(
            vertical="top", wrap_text=True)

    line("Prepared for", client_name)
    line("Prepared on", datetime.now().strftime("%d %b %Y"))
    line("Requirement", query)
    if criteria.get("product_note"):
        line("Product", criteria["product_note"])
    line("Options included", len(records))

    markets = []
    for record in records:
        market = _market_of(record)
        if market not in markets:
            markets.append(market)
    if markets:
        line("Micro-markets", ", ".join("%s (%d)" % (
            m, sum(1 for r in records if _market_of(r) == m)) for m in markets))

    landlords = sorted({r.get("developer_landlord") for r in records
                        if r.get("developer_landlord")})
    if landlords:
        line("Landlords / operators", ", ".join(landlords))

    ws.append([])
    ws.append(["Each micro-market has its own sheet, laid out as the supply workbook is. "
               "Every figure is taken from the database as at the date above."])
    ws.cell(row=ws.max_row, column=1).font = Font(size=9, italic=True, color="7F7F7F")
